"""Build a case-study doc (docs 3a–3f): one real case (2015+) per syllabus case-study slot of an H2 topic.

Content: tools/data/case_studies_<id>.json. Charts: tools/case_charts.py -> tools/data/case_charts/.
usage: build_case_studies.py <content.json> <outdir> <version>
"""
import json, os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from build_docs import base_doc, cell_text, shade, repeat_header, keep_row_together, GREY
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
CHARTS = os.path.join(HERE, 'data', 'case_charts')
LINK = RGBColor(0x1F, 0x4E, 0x99)


def hyperlink(par, text, url, size=7.5):
    rid = par.part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',
                             is_external=True)
    h = OxmlElement('w:hyperlink'); h.set(qn('r:id'), rid)
    r = OxmlElement('w:r'); rpr = OxmlElement('w:rPr')
    for tag, val in (('w:color', '1F4E99'), ('w:u', 'single'), ('w:sz', str(int(size * 2)))):
        el = OxmlElement(tag); el.set(qn('w:val'), val); rpr.append(el)
    r.append(rpr); t = OxmlElement('w:t'); t.text = text; r.append(t); h.append(r)
    par._p.append(h)


def small(par, text, size=9, italic=False, bold=False, colour=None):
    r = par.add_run(text)
    r.font.size, r.italic, r.bold = Pt(size), italic, bold
    if colour is not None:
        r.font.color.rgb = colour
    return r


def tight(par, before=0, after=2):
    par.paragraph_format.space_before, par.paragraph_format.space_after = Pt(before), Pt(after)


def set_widths(t, widths):
    """Fixed column widths (Word reads cell widths, LibreOffice the grid)."""
    t.autofit = False
    for gc, w in zip(t._tbl.tblGrid.findall(qn('w:gridCol')), widths):
        gc.set(qn('w:w'), str(int(Cm(w).twips)))
    for row in t.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = Cm(w)


def plan_table(doc, rows):
    cols = ['#', 'Sub-topic', 'Type', 'Syllabus point', 'Case study']
    t = doc.add_table(rows=1, cols=len(cols)); t.style = 'Table Grid'
    for c, h in zip(t.rows[0].cells, cols):
        cell_text(c, [(h, True, False)], size=8); shade(c, 'D9E2F3')
    repeat_header(t.rows[0])
    for row in rows:
        r = t.add_row(); keep_row_together(r)
        done = row[4] != 'v1'
        for c, v in zip(r.cells, row[:4] + [row[4] if done else 'To come in v1']):
            cell_text(c, [(v, False, not done and v == 'To come in v1')], size=8)
            if done:
                shade(c, 'E2EFDA')
    set_widths(t, (0.8, 3.4, 2.4, 6.4, 5.0))


def fact_table(doc, facts):
    t = doc.add_table(rows=0, cols=2); t.style = 'Table Grid'
    for k, v in facts:
        r = t.add_row(); keep_row_together(r)
        cell_text(r.cells[0], [(k, True, False)], size=8.5); shade(r.cells[0], 'F2F2F2')
        cell_text(r.cells[1], [(v, False, False)], size=8.5)
    set_widths(t, (2.6, 15.4))


def add_case(doc, c):
    doc.add_page_break()
    h = doc.add_heading(c['heading'], level=1); tight(h, 0, 2)
    p = doc.add_paragraph(); tight(p, 0, 2)
    small(p, 'Case study: ', 11, bold=True); small(p, c['case'], 11, bold=True, colour=LINK)
    p = doc.add_paragraph(); tight(p, 0, 4)
    small(p, 'Syllabus: ' + c['syllabus'], 8.5, italic=True, colour=GREY)
    fact_table(doc, c['facts'])
    p = doc.add_paragraph(); tight(p, 4, 0); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(os.path.join(CHARTS, c['image']), width=Cm(15.5))
    p = doc.add_paragraph(); tight(p, 0, 4); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    small(p, c['caption'], 8, italic=True, colour=GREY)
    for s in c['sections']:
        p = doc.add_paragraph(); tight(p, 3, 1); small(p, s['title'], 10, bold=True)
        if s.get('text'):
            p = doc.add_paragraph(); tight(p, 0, 2); small(p, s['text'], 9.5)
        for b in s.get('bullets', []):
            p = doc.add_paragraph(style='List Bullet'); tight(p, 0, 1); small(p, b, 9.5)
    p = doc.add_paragraph(); tight(p, 4, 0); small(p, 'Sources', 8, bold=True, colour=GREY)
    for i, (label, url) in enumerate(c['sources'], 1):
        p = doc.add_paragraph(); tight(p, 0, 0)
        small(p, f'{i}. ', 7.5, colour=GREY); hyperlink(p, label, url)


def build(spec, outdir, version):
    doc = base_doc(spec['title'], version)
    for para in spec['intro']:
        p = doc.add_paragraph(); tight(p, 0, 4); small(p, para, 10)
    doc.add_heading(spec['plan_title'], level=2)
    plan_table(doc, spec['plan'])
    p = doc.add_paragraph(); tight(p, 3, 0)
    small(p, 'Green rows are in this version. Each case starts on a new page.', 8, italic=True, colour=GREY)
    for c in spec['cases']:
        add_case(doc, c)
    path = os.path.join(outdir, f"{spec['file']} {version}.docx")
    doc.save(path)
    return path


if __name__ == '__main__':
    print(build(json.load(open(sys.argv[1], encoding='utf8')), sys.argv[2], sys.argv[3]))
