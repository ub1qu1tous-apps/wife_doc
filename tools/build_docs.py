"""Build the consolidated Word docs from questions.json.

Doc 1: topic grid (H1 table, H2 table): rows = school + year, columns = question slots,
       cells = syllabus topic(s) of each question.
Doc 2: questions listed by syllabus topic, split into 6 docs
       (H1/H2 x fieldwork, structured, essay); each entry = school + year,
       full question text, then the resources (diagrams) it uses.

usage: build_docs.py <questions.json> <outdir> <version> [--preview] [--only grid|review|bytopic] [--doc 2a]
"""
import io, json, os, sys
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
TOPICS = json.load(open(os.path.join(HERE, 'syllabus_topics.json')))
GREY = RGBColor(0x59, 0x59, 0x59)

H1_SLOTS = [('Q1', 'Sec A', 'Fieldwork'), ('Q2', 'Sec B', 'Structured'), ('Q3', 'Sec B', 'Structured'),
            ('Q4', 'Sec C', 'Essay'), ('Q5', 'Sec C', 'Essay'), ('Q6', 'Sec C', 'Essay')]
H2_SLOTS = [('P1 Q1', 'Cluster 1', 'Structured'), ('P1 Q2', 'Cluster 2', 'Structured'),
            ('P1 Q3', 'Cluster 1', 'Essay'), ('P1 Q4', 'Cluster 1', 'Essay'),
            ('P1 Q5', 'Cluster 2', 'Essay'), ('P1 Q6', 'Cluster 2', 'Essay'),
            ('P2 Q1', 'Fieldwork', 'Structured'), ('P2 Q2', 'Cluster 3', 'Structured'),
            ('P2 Q3', 'Cluster 3', 'Essay'), ('P2 Q4', 'Cluster 3', 'Essay')]


def topic_label(level, code):
    """'1.2.2' -> '1.2.2 Possible Effects of Climate Change'; key question '1.2.2b' -> '1.2.2b Possible Effects of ... – <KQ>'."""
    t = TOPICS[level]
    if code in t.get('kq', {}):
        return f"{code} {t['subtopics'][code[:-1]]} – {t['kq'][code]}"
    return f"{code} {t['subtopics'].get(code) or t['topics'].get(code) or code}"


def short_label(level, code):
    """'1.2.2' -> '1.2.2 Effects of CC'; key-question code '1.1.1b' -> '1.1.1b SUD: Urban population trends' (grid cells)."""
    t = TOPICS[level]
    if code in t.get('kq', {}):
        return f"{code} {t['short'][code[:-1]]}: {t['kq'][code]}"
    return f"{code} {t['short'].get(code) or t['topics'].get(code) or code}"


def base_doc(title, version, landscape=False, preview=False):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    if landscape:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for side in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
        setattr(sec, side, Cm(1.5))
    st = doc.styles['Normal']
    st.font.name, st.font.size = 'Arial', Pt(10)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
    for name, size in (('Title', 18), ('Heading 1', 14), ('Heading 2', 12), ('Heading 3', 11)):
        s = doc.styles[name]
        s.font.name, s.font.size = 'Arial', Pt(size)
        s.element.rPr.rFonts.set(qn('w:asciiTheme'), 'Arial') if False else None
    doc.add_heading(f"{title} — {version}" + (" (PREVIEW)" if preview else ''), level=0)
    p = doc.add_paragraph()
    r = p.add_run("Topics follow the SEAB syllabus: H1 Geography 8834 (2026) and H2 Geography 9173 (2026).")
    r.italic, r.font.color.rgb, r.font.size = True, GREY, Pt(9)
    return doc


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hex_fill)
    tcPr.append(shd)


def keep_row_together(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement('w:cantSplit'); el.set(qn('w:val'), 'true'); trPr.append(el)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement('w:tblHeader'); el.set(qn('w:val'), 'true'); trPr.append(el)


def cell_text(cell, lines, size=8):
    cell.text = ''
    first = True
    for text, bold, grey in lines:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(text)
        r.bold, r.font.size = bold, Pt(size)
        if grey:
            r.italic, r.font.color.rgb = True, GREY


def grid_table(doc, level, slots, qs, schools_years, title, missing=()):
    doc.add_heading(title, level=1)
    t = doc.add_table(rows=2, cols=len(slots) + 1)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr, sub = t.rows[0], t.rows[1]
    cell_text(hdr.cells[0], [('School & Year', True, False)])
    cell_text(sub.cells[0], [('', False, False)])
    for i, (slot, sec, kind) in enumerate(slots, 1):
        cell_text(hdr.cells[i], [(slot, True, False)])
        cell_text(sub.cells[i], [(f'{sec} · {kind}', False, True)], size=7)
    for row in (hdr, sub):
        repeat_header(row)
        for c in row.cells:
            shade(c, 'D9E2F3')
    for school, year in schools_years:
        row = t.add_row()
        keep_row_together(row)
        cells = row.cells
        cell_text(cells[0], [(f'{school} {year}', True, False)])
        for i, (slot, _, _) in enumerate(slots, 1):
            q = qs.get((school, year, slot))
            if not q:
                cell_text(cells[i], [('—', False, True)])
                continue
            lines = [(short_label(level, c), True, False) for c in q['topics']]
            if q.get('fieldwork_area'):
                lines = [(f"Fieldwork – {q['fieldwork_area']}", True, False)]
            cell_text(cells[i], lines)
    if missing:
        r = doc.add_paragraph().add_run('No paper in sources: ' + ', '.join(missing))
        r.italic, r.font.size, r.font.color.rgb = True, Pt(8), GREY
    widths = [Cm(2.2)] + [Cm((doc.sections[0].page_width.cm - 3.0 - 2.2) / len(slots))] * len(slots)
    for row in t.rows:
        for w, c in zip(widths, row.cells):
            c.width = w


def build_grid(questions, outdir, version, preview):
    doc = base_doc('Other JC Prelims – Topics by School (H1 & H2)', version, landscape=True, preview=preview)
    for level, slots, title in (('H1', H1_SLOTS, 'H1 Geography (8834) – Paper 1'),
                                ('H2', H2_SLOTS, 'H2 Geography (9173) – Papers 1 & 2')):
        qs = {(q['school'], q['year'], q['slot']): q for q in questions if q['level'] == level}
        rows = sorted({(s, y) for s, y, _ in qs})
        all_rows = sorted({(q['school'], q['year']) for q in questions})
        missing = [f'{s} {y}' + (' (password-protected QP, left out)' if (s, y) == ('VJC', 2026) else '')
                   for s, y in all_rows if (s, y) not in rows]
        grid_table(doc, level, slots, qs, rows, title, missing)
    doc.add_heading('Syllabus topic codes', level=1)
    for level, name in (('H1', 'H1 Geography 8834'), ('H2', 'H2 Geography 9173')):
        p = doc.add_paragraph()
        p.add_run(name + ': ').bold = True
        tops = TOPICS[level]
        items = [f"{c} {tops['subtopics'][c]}" + (' — ' + ', '.join(f"{k[-1]}) {v}" for k, v in tops['kq'].items() if k[:-1] == c) if 'kq' in tops else '')
                 for c in tops['subtopics']]
        items.append('Fieldwork – Elderly needs / Community response to CC' + (' / Fluvial flood risk' if level == 'H2' else '') + ' / Others')
        r = p.add_run('; '.join(items))
        r.font.size = Pt(8)
    name = '1 Topic Grid by School (H1 & H2)'
    path = os.path.join(outdir, f"{name} {version}{' PREVIEW' if preview else ''}.docx")
    doc.save(path)
    return path


def compact_image(path, max_px=1300):
    """Shrink a crop for embedding: cap its size, then keep whichever of JPEG / 256-colour PNG is smaller."""
    im = (path if isinstance(path, Image.Image) else Image.open(path)).convert('RGB')
    im.thumbnail((max_px, max_px * 3))
    jpg, png = io.BytesIO(), io.BytesIO()
    im.save(jpg, 'JPEG', quality=82, optimize=True)
    im.quantize(256).save(png, 'PNG', optimize=True)
    best = min(jpg, png, key=lambda b: b.getbuffer().nbytes)
    best.seek(0)
    return best


def split_tall(im, page_h):
    """Cut a very tall image into page-high pieces, each cut at the whitest row in the last fifth of the piece."""
    g = im.convert('L')
    pieces, y0 = [], 0
    while im.height - y0 > page_h:
        lo, hi = y0 + int(page_h * 0.8), y0 + page_h
        cut = max(range(lo, hi), key=lambda y: (sum(g.crop((0, y, g.width, y + 1)).getdata()), -abs(y - hi)))
        pieces.append(im.crop((0, y0, im.width, cut)))
        y0 = cut
    pieces.append(im.crop((0, y0, im.width, im.height)))
    return pieces


def add_question(doc, q, img_root, new_page=False):
    level = q['level']
    h = doc.add_heading(f"{q['school']} {q['year']} — {q['paper']}, {q['qno']}  [{q['marks']} marks]", level=3)
    h.paragraph_format.page_break_before = new_page
    p = doc.add_paragraph()
    r = p.add_run('Syllabus: ' + '; '.join(topic_label(level, c) for c in q['topics'])
                  + (f"   |   Fieldwork area: {q['fieldwork_area']}" if q.get('fieldwork_area') else ''))
    r.italic, r.font.size, r.font.color.rgb = True, Pt(8.5), GREY
    section, *stem = q['stem']
    r = doc.add_paragraph().add_run(section)
    r.font.size, r.font.color.rgb, r.bold = Pt(8.5), GREY, True
    for para in stem:
        doc.add_paragraph(para).paragraph_format.space_after = Pt(4)
    right = doc.sections[0].page_width - doc.sections[0].left_margin - doc.sections[0].right_margin
    for label, text, marks in q.get('parts', []):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent, pf.first_line_indent, pf.space_after = Cm(1.0), Cm(-1.0), Pt(4)
        pf.tab_stops.add_tab_stop(Cm(1.0))
        pf.tab_stops.add_tab_stop(right, WD_TAB_ALIGNMENT.RIGHT)
        p.add_run(f'{label}\t{text}\t[{marks}]')
    for res in q.get('resources', []):
        doc.add_paragraph()
        path = os.path.join(img_root, res)
        w, h = Image.open(path).size
        if h / w > 1.9:   # would shrink below ~12 cm wide on one page (multi-page questionnaires): split at blank rows
            for piece in split_tall(Image.open(path).convert('RGB'), int(w * 23.5 / 15.5)):
                doc.add_picture(compact_image(piece), width=Cm(15.5))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            continue
        if h / w * 15.5 > 23.5:     # very tall resource: fit to page height instead
            doc.add_picture(compact_image(path), height=Cm(23.5))
        else:
            doc.add_picture(compact_image(path), width=Cm(15.5))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


def build_by_topic(questions, outdir, version, preview, img_root):
    docs = [('H1', 'fieldwork', '2a H1 Fieldwork Questions by Topic', 'H1 Fieldwork Questions by Topic (8834 Sec A)'),
            ('H1', 'structured', '2b H1 Structured Questions by Topic', 'H1 Structured Questions by Topic (8834 Sec B)'),
            ('H1', 'essay', '2c H1 Essay Questions by Topic', 'H1 Essay Questions by Topic (8834 Sec C)'),
            ('H2', 'fieldwork', '2d H2 Fieldwork Questions by Topic', 'H2 Fieldwork Questions by Topic (9173 P2 Sec A)'),
            ('H2', 'structured', '2e H2 Structured Questions by Topic', 'H2 Structured Questions by Topic (9173 P1 Sec A, P2 Sec B)'),
            ('H2', 'essay', '2f H2 Essay Questions by Topic', 'H2 Essay Questions by Topic (9173 P1 Sec B, P2 Sec C)')]
    paths = []
    only_doc = sys.argv[sys.argv.index('--doc') + 1] if '--doc' in sys.argv else None
    for level, kind, fname, title in docs:
        if only_doc and not fname.startswith(only_doc + ' '):
            continue
        doc = base_doc(title, version, preview=preview)
        sel = [q for q in questions if q['level'] == level and q['kind'] == kind and q.get('stem')]
        tops = TOPICS[level]
        # group under primary topic (structured) or primary sub-topic (essay)
        def key(q):
            return q['primary'] + ('|' + q['fieldwork_area'] if q.get('fieldwork_area') else '')

        def order(g):
            code, _, area = g.partition('|')
            return [int(x) for x in code.split('.')] + [area == 'Others', area]

        def heading(g):
            code, _, area = g.partition('|')
            return f'Fieldwork – {area}' if area else topic_label(level, code)
        # each school's question(s) start on a new page for printing; a heading stays on the page of its first question
        started = fresh = False
        for cluster, cl_name in tops['clusters'].items():
            in_cluster = [q for q in sel if q['primary'].split('.')[0] == cluster]
            if not in_cluster:
                continue
            h = doc.add_heading(f'Cluster {cluster}: {cl_name}', level=1)
            h.paragraph_format.page_break_before, fresh = started, started
            for g in sorted({key(q) for q in in_cluster}, key=order):
                h = doc.add_heading(heading(g), level=2)
                if started and not fresh:
                    h.paragraph_format.page_break_before = fresh = True
                prev = None
                for q in sorted((q for q in in_cluster if key(q) == g), key=lambda q: (q['school'], q['year'])):
                    add_question(doc, q, img_root, new_page=started and not fresh and (q['school'], q['year']) != prev)
                    started, fresh, prev = True, False, (q['school'], q['year'])
        path = os.path.join(outdir, f"{fname} {version}{' PREVIEW' if preview else ''}.docx")
        doc.save(path)
        paths.append(path)
    return paths


def build_review(review, outdir, version, preview):
    """Doc 0: things the user should check (Others fieldwork, excluded files, paper errors)."""
    doc = base_doc('Items to Review', version, preview=preview)
    for sec in review:
        doc.add_heading(sec['title'], level=1)
        if sec.get('note'):
            r = doc.add_paragraph().add_run(sec['note'])
            r.italic, r.font.size, r.font.color.rgb = True, Pt(9), GREY
        t = doc.add_table(rows=1, cols=len(sec['columns']))
        t.style = 'Table Grid'
        for c, h in zip(t.rows[0].cells, sec['columns']):
            cell_text(c, [(h, True, False)], size=9)
            shade(c, 'D9E2F3')
        repeat_header(t.rows[0])
        for row in sec['rows']:
            for c, v in zip(t.add_row().cells, row):
                cell_text(c, [(str(v), False, False)], size=9)
    path = os.path.join(outdir, f"0 Items to Review {version}{' PREVIEW' if preview else ''}.docx")
    doc.save(path)
    return path


if __name__ == '__main__':
    qfile, outdir, version = sys.argv[1:4]
    preview = '--preview' in sys.argv
    questions = json.load(open(qfile))
    os.makedirs(outdir, exist_ok=True)
    img_root = os.path.dirname(os.path.abspath(qfile))
    only = sys.argv[sys.argv.index('--only') + 1] if '--only' in sys.argv else None
    review_file = os.path.join(img_root, 'review.json')
    if only in (None, 'review') and os.path.exists(review_file):
        print(build_review(json.load(open(review_file)), outdir, version, preview))
    if only in (None, 'grid'):
        print(build_grid(questions, outdir, version, preview))
    if only in (None, 'bytopic'):
        for p in build_by_topic(questions, outdir, version, preview, img_root):
            print(p)
